"""The Tally must be live formulas that agree with the solver's counts.

We emulate COUNTIF over the student-sheet columns (the spreadsheet's single
source of truth) and confirm it equals what the solver assigned — so the live
tally and the generated schedule can never silently disagree.
"""

from collections import Counter
from pathlib import Path

import pytest
from openpyxl import load_workbook

from cjai.pipeline import GRADES, run_grade
from cjai.report import GradeRun, _slot_col, write_output

WORKBOOK = Path(__file__).resolve().parent.parent / "data" / "scheduling-26-27.xlsx"

pytestmark = pytest.mark.skipif(not WORKBOOK.exists(), reason="workbook not present")


def _col_counts(ws, col: int) -> Counter:
    c: Counter = Counter()
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, col).value
        if isinstance(v, (int, float)):
            c[int(v)] += 1
    return c


def test_tally_is_live_formulas_matching_solver(tmp_path):
    problem, result, anomalies, _ = run_grade(str(WORKBOOK), 6)
    cfg = GRADES[6]
    out = tmp_path / "out.xlsx"
    write_output(str(WORKBOOK), str(out), [GradeRun(cfg, problem, result, anomalies)])

    wb = load_workbook(out)  # keep formulas
    tally = wb["Tally"]
    assert str(tally.cell(2, 7).value).startswith("=COUNTIF")  # assigned: live count
    # capacity: live reference into the offerings sheet, not a static number
    capacity_formula = str(tally.cell(2, 8).value)
    assert capacity_formula.startswith("=")
    assert cfg.offerings_sheet in capacity_formula

    nm = wb[cfg.nonmusic_sheet]
    mu = wb[cfg.music_sheet]
    cache: dict = {}

    def counts(ws, col):
        return cache.setdefault((ws.title, col), _col_counts(ws, col))

    for s in problem.sections:
        nm_col = _slot_col(cfg, s.slot, cfg.nonmusic_col0)
        mu_col = _slot_col(cfg, s.slot, cfg.music_col0)
        emulated = counts(nm, nm_col).get(s.course_no, 0) + counts(mu, mu_col).get(s.course_no, 0)
        assert emulated == result.load.get(s.key, 0), f"{s.course_name} {s.slot}"
