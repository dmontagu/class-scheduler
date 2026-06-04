"""Checks for the PowerSchool submission sheet."""

from pathlib import Path

import pytest
from openpyxl import load_workbook

from cjai.pipeline import GRADES, run_grade
from cjai.report import GradeRun, _ordered_electives, write_output

WORKBOOK = Path(__file__).resolve().parent.parent / "data" / "scheduling-26-27.xlsx"

pytestmark = pytest.mark.skipif(not WORKBOOK.exists(), reason="workbook not present")


def test_ordered_electives_distinct_and_complete():
    problem, _, _, _ = run_grade(str(WORKBOOK), 6)
    cfg = GRADES[6]
    for s in problem.students:
        electives = _ordered_electives(cfg, s)
        assert len(electives) == len(s.open_slots)
        assert len(electives) == len(set(electives))  # no repeats


def test_submission_sheet_shape(tmp_path):
    problem, result, anomalies, _ = run_grade(str(WORKBOOK), 6)
    out = tmp_path / "out.xlsx"
    write_output(str(WORKBOOK), str(out), [GradeRun(GRADES[6], problem, result, anomalies)])

    ws = load_workbook(out, data_only=True)["PowerSchool Submission"]
    assert ws.cell(1, 1).value == "Student Numbers"
    assert ws.cell(1, 2).value == "E1"
    # one row per student that needs electives (double-music excluded)
    needing = [s for s in problem.students if s.open_slots]
    data_rows = sum(1 for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value)
    assert data_rows == len(needing)
