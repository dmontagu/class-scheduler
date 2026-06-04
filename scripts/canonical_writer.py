# ruff: noqa: E501  (long literal strings: schema + README prose read better unwrapped)
"""Shared writer for the canonical scheduling data-model workbooks (PRD §9.5).

One source of truth for the per-grade tidy-table schema, used by both the mock
example generator and the real-data extractor. PRIVACY: the `students` sheet
holds IDs only — no student names — so nothing sensitive is passed around or
sent to an LLM. (Course/teacher names are not student PII and are kept.)
"""

from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

CATALOG_TAB = "1F6FEB"  # blue
STUDENT_TAB = "1A7F37"  # green
RULES_TAB = "9A6700"  # amber
README_TAB = "6E7781"  # grey

HEADER_FILL = PatternFill("solid", fgColor="DDE7F5")
HEADER_FONT = Font(bold=True)

CATEGORIES = [
    "Art", "Tech Ed", "FCS", "Language", "Performing Arts", "Music",
    "Core-Math", "Core-ELA/SS", "Core-Science", "Core-Health", "Global",
]
TERM_TYPES = ["year_long", "semester", "quarter"]
GRADES = ["6", "7", "8"]
YESNO = ["yes", "no"]

# canonical schema — one entry per sheet, in file order
SHEETS: dict[str, dict] = {
    "courses": dict(tab=CATALOG_TAB,
                    headers=["course_id", "name", "category", "term_type", "grade", "notes"],
                    dropdowns={"category": CATEGORIES, "term_type": TERM_TYPES, "grade": GRADES},
                    widths={"name": 22, "category": 16, "term_type": 12, "notes": 50}),
    "sections": dict(tab=CATALOG_TAB,
                     headers=["section_id", "course_id", "course_name", "teacher", "block",
                              "runs_quarters", "capacity_target", "capacity_max"],
                     widths={"section_id": 14, "course_name": 20, "teacher": 12,
                             "runs_quarters": 14, "capacity_target": 15, "capacity_max": 13}),
    "blocks": dict(tab=CATALOG_TAB, headers=["grade", "block_label", "description"],
                   dropdowns={"grade": GRADES}, widths={"block_label": 12, "description": 66}),
    "students": dict(tab=STUDENT_TAB, headers=["student_id", "grade"],
                     dropdowns={"grade": GRADES}, widths={"student_id": 12}),
    "enrollments_music": dict(tab=STUDENT_TAB,
                              headers=["student_id", "course_id", "section_id", "term", "source", "locked", "notes"],
                              dropdowns={"locked": YESNO},
                              widths={"section_id": 14, "term": 8, "source": 16, "locked": 8, "notes": 58}),
    "enrollments_language": dict(tab=STUDENT_TAB,
                                 headers=["student_id", "course_id", "section_id", "term", "source", "locked", "notes"],
                                 dropdowns={"locked": YESNO},
                                 widths={"section_id": 14, "term": 8, "source": 16, "locked": 8, "notes": 58}),
    "enrollments_sped": dict(tab=STUDENT_TAB,
                             headers=["student_id", "course_id", "section_id", "term", "source", "locked", "notes"],
                             dropdowns={"locked": YESNO},
                             widths={"section_id": 14, "term": 8, "source": 16, "locked": 8, "notes": 58}),
    "track_members": dict(tab=STUDENT_TAB, headers=["track_id", "student_id", "notes"],
                          widths={"track_id": 18, "notes": 40}),
    "track_sections": dict(tab=STUDENT_TAB, headers=["track_id", "section_id", "notes"],
                           widths={"track_id": 18, "section_id": 14, "notes": 36}),
    "incompatibilities": dict(tab=STUDENT_TAB,
                              headers=["student_id_a", "student_id_b", "hard", "reason"],
                              dropdowns={"hard": YESNO},
                              widths={"student_id_a": 13, "student_id_b": 13, "reason": 46}),
    "pairings": dict(tab=STUDENT_TAB,
                     headers=["student_id_a", "student_id_b", "min_together", "max_together", "hard", "reason"],
                     dropdowns={"hard": YESNO},
                     widths={"student_id_a": 13, "student_id_b": 13, "min_together": 13,
                             "max_together": 13, "reason": 56}),
    "preferences (optional)": dict(tab=STUDENT_TAB, headers=["student_id", "course_id", "rank", "notes"],
                                   widths={"notes": 52}),
    "rules": dict(tab=RULES_TAB, headers=["rule_id", "scope", "value", "hard", "description"],
                  dropdowns={"hard": YESNO},
                  widths={"rule_id": 30, "scope": 14, "value": 8, "description": 72}),
}


def _quoted(options: list[str]) -> str:
    return '"' + ",".join(options) + '"'


def new_workbook() -> Workbook:
    wb = Workbook()
    default = wb.active
    if default is not None:
        wb.remove(default)
    return wb


def add(wb: Workbook, sheet: str, rows: list[list]):
    """Write a known canonical sheet using its registered schema."""
    spec = SHEETS[sheet]
    headers = spec["headers"]
    ws = wb.create_sheet(sheet)
    ws.sheet_properties.tabColor = spec["tab"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    for r, row in enumerate(rows, 2):
        for c, val in enumerate(row, 1):
            ws.cell(r, c, val)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    widths = spec.get("widths", {})
    for c, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(c)].width = widths.get(h, max(12, len(h) + 3))
    for hname, options in spec.get("dropdowns", {}).items():
        col = get_column_letter(headers.index(hname) + 1)
        dv = DataValidation(type="list", formula1=_quoted(options), allow_blank=True)
        dv.add(f"{col}2:{col}1000")
        ws.add_data_validation(dv)
    return ws


def default_rules() -> list[list]:
    """Suggested global policy rows (same for every grade; team can tune)."""
    return [
        ["no_pair_shares_all_classes", "all students", "", "no", "No two students should have EVERY class together (variety / social mixing)"],
        ["max_shared_classes_any_pair", "all students", 3, "no", "Cap on how many classes any two students share (tune for variety)"],
        ["no_same_category_same_quarter", "all students", "", "no", "Avoid two electives of one category in the same quarter (near-hard)"],
        ["no_two_languages_same_quarter", "grades 6,7", "", "yes", "A student may not take two Language-category electives in one quarter (6th/7th)"],
        ["avoid_teacher_back_to_back", "all students", "", "no", "Avoid the same teacher in consecutive quarters"],
        ["maximize_category_variety", "all students", "", "no", "Prefer distinct categories; weight music students (fewer slots) higher"],
        ["section_target_size", "all sections", 30, "no", "Aim for this size; hard cap is each section's capacity_max"],
    ]


def add_readme(wb: Workbook, grade: int, *, real: bool):
    ws = wb.create_sheet("READ ME")
    ws.sheet_properties.tabColor = README_TAB
    ws.column_dimensions["A"].width = 114
    H, T = True, False
    kind = ("REAL starter data (extracted from this year's workbook where available; "
            "blank sheets are for you to fill in)") if real else "EXAMPLE with mock data"
    lines = [
        (f"Canonical Scheduling Data Model — GRADE {grade} — {kind}", H),
        ("", T),
        ("ONE WORKBOOK PER GRADE (mixing grades in one file is too much). This is the GRADE", T),
        (f"{grade} file; all data here is grade {grade}.", T),
        ("", T),
        ("PRIVACY: the 'students' sheet lists student IDs only — NO NAMES. Student names are", H),
        ("sensitive, so we keep them out of the data the tool handles (and never send them to an", T),
        ("LLM). To see who an ID is, look it up in PowerSchool. (Course and teacher names are not", T),
        ("student data and are kept.)", T),
        ("", T),
        ("The big idea: one fact per row ('tidy' tables) instead of one wide visual grid — easy to", T),
        ("add to as info trickles in, easy to validate, easy to combine (the tool joins on ID columns).", T),
        ("", T),
        ("Sheet types (tab colors):", H),
        ("  • CATALOG (blue): courses, sections, blocks — slow-changing, owned by the scheduling lead.", T),
        ("  • PER-STUDENT (green): students; enrollments_music / _language / _sped; track_members /", T),
        ("      track_sections; incompatibilities; pairings; preferences — fast-changing, many sources.", T),
        ("  • RULES (amber): global policies that apply broadly, not to one student.", T),
        ("", T),
        ("Enrollments are SPLIT by source so different people maintain different sheets:", H),
        ("  enrollments_music    <- music sign-up (Band/Choir/Orchestra/Music Theory).", T),
        ("  enrollments_language <- world-language sign-up (year-long languages).", T),
        ("  enrollments_sped     <- SpEd case managers (hand-built, locked placements).", T),
        ("  The solver only DECIDES electives; everything here is collected and honored exactly.", T),
        ("", T),
        ("Relationship sheets (same shape, opposite meaning):", H),
        ("  incompatibilities = must NOT share a section ('hard'=yes never violated).", T),
        ("  pairings          = SHOULD share classes: min_together (e.g. >=2) and max_together (so they", T),
        ("      don't share EVERY class). 'hard'=yes required; 'no' = strong preference. See also the", T),
        ("      'no_pair_shares_all_classes' rule, which applies to everyone.", T),
        ("", T),
        ("Key columns: *_id are stable join keys; term_type = year_long(4q)/semester(2q)/quarter(1q);", T),
        ("locked 'yes' = never move; hard 'yes' = never violate, 'no' = soft (optimized, can trade off).", T),
    ]
    if real:
        lines += [
            ("", T),
            ("WHAT'S PRE-FILLED vs BLANK in this starter file:", H),
            ("  Pre-filled from data: courses, sections (electives + music where known), students,", T),
            ("      enrollments_music, enrollments_language. Some teacher/capacity cells for music/core", T),
            ("      sections are blank where the source didn't have them — please confirm.", T),
            ("  BLANK for you to fill: enrollments_sped, track_members, track_sections, incompatibilities,", T),
            ("      pairings, preferences. 'rules' is pre-filled with suggested defaults to tune.", T),
            ("  Please review category and term_type on 'courses' — best-guessed, may need correcting.", T),
        ]
    for r, (text, is_head) in enumerate(lines, 1):
        cell = ws.cell(r, 1, text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if is_head:
            cell.font = Font(bold=True, size=13 if r == 1 else 11)
    ws.sheet_view.showGridLines = False
    return ws
